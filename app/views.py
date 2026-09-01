import os
import base64
import json
import uuid
import threading
import numpy as np
import cv2 as cv
from datetime import timedelta, datetime
from django.utils import timezone
from django.conf import settings
from django.core.cache import cache
from django.http import StreamingHttpResponse, HttpResponse
from django.shortcuts import get_object_or_404
from django.db.models.functions import TruncDate
from django.db.models import Count, Min, Max
from openpyxl import Workbook
from rest_framework import viewsets, permissions, status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.decorators import api_view, permission_classes
from .models import User, Person, PersonImage, PersonEmbedding, RecognitionLog, Camera, Notification, DevicePushToken
from .serializers import (
    UserSerializer,
    PersonSerializer,
    PersonImageSerializer,
    RecognitionLogSerializer,
    CameraSerializer,
    NotificationSerializer
)
from app.utils.video_processor import (
    gen_face_login_frames,
    gen_frames,
    get_yolo_model,
    rotate_and_flip_frame,
    draw_detection_box,
    _feed_stats,
    _face_login_verified
)
from app.utils.alerts import send_alerts
from app.utils.reports import _fill_excel_worksheet, _fill_summary_sheet, _fill_camera_analytics_sheet


# ==============================================================================
# REST FRAMEWORK USER & AUTH API VIEWS
# ==============================================================================

class UserViewSet(viewsets.ModelViewSet):
    serializer_class = UserSerializer
    permission_classes = [permissions.IsAdminUser]

    def get_queryset(self):
        return User.objects.all().order_by('id')

class UserMeView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        serializer = UserSerializer(request.user)
        return Response(serializer.data)


from django.contrib.auth.tokens import PasswordResetTokenGenerator
from django.utils.encoding import force_str
from django.utils.http import urlsafe_base64_decode

_password_reset_token_generator = PasswordResetTokenGenerator()


class ForgotPasswordAPIView(APIView):
    permission_classes = [permissions.AllowAny]

    def post(self, request):
        username = request.data.get('username')
        code = request.data.get('code')
        new_password = request.data.get('new_password')
        token = request.data.get('token')
        uid = request.data.get('uid')

        if not new_password:
            return Response({'error': 'New password is required.'}, status=status.HTTP_400_BAD_REQUEST)

        # 1. Django Secure Token-based reset (Fix 9)
        if token and uid:
            try:
                user_id = force_str(urlsafe_base64_decode(uid))
                user = User.objects.get(pk=user_id)
                if _password_reset_token_generator.check_token(user, token):
                    user.set_password(new_password)
                    user.save()
                    return Response({'message': 'Password reset successfully. You can now login.'}, status=status.HTTP_200_OK)
                return Response({'error': 'Invalid or expired password reset token.'}, status=status.HTTP_400_BAD_REQUEST)
            except Exception:
                return Response({'error': 'Invalid reset link or parameters.'}, status=status.HTTP_400_BAD_REQUEST)

        # 2. 2-Step system recovery code
        if not username or not code:
            return Response({'error': 'Username and recovery code (or reset token) are required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            user = User.objects.get(username=username, code=code)
            user.set_password(new_password)
            user.save()
            return Response({'message': 'Password reset successfully. You can now login.'}, status=status.HTTP_200_OK)
        except User.DoesNotExist:
            return Response({'error': 'Invalid username or recovery code.'}, status=status.HTTP_400_BAD_REQUEST)


# ==============================================================================
# REST FRAMEWORK CAMERA VIEWSET
# ==============================================================================

class CameraViewSet(viewsets.ModelViewSet):
    queryset = Camera.objects.all().order_by('-created_at')
    serializer_class = CameraSerializer
    permission_classes = [permissions.IsAuthenticated]


# ==============================================================================
# REST FRAMEWORK PERSON & DATASET VIEWS
# ==============================================================================

class PersonViewSet(viewsets.ModelViewSet):
    queryset = Person.objects.all().order_by('name')
    serializer_class = PersonSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context['request'] = self.request
        return context


class PersonImageDeleteView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def delete(self, request, image_id):
        img = get_object_or_404(PersonImage, id=image_id)
        if img.image and os.path.isfile(img.image.path):
            try:
                os.remove(img.image.path)
            except Exception:
                pass
        img.delete()
        return Response({'status': 'success', 'message': 'Image deleted'}, status=status.HTTP_204_NO_CONTENT)


class DatasetUploadView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        person_name = request.data.get('name')
        images = request.FILES.getlist('images')

        if not person_name:
            return Response({'error': 'Person name is required'}, status=status.HTTP_400_BAD_REQUEST)

        person, created = Person.objects.get_or_create(name=person_name.strip())

        uploaded_images = []
        embedding_count = 0
        if images:
            for img in images:
                pi = PersonImage.objects.create(person=person, image=img)
                uploaded_images.append(pi)

                # Auto-compute ArcFace embedding for the uploaded image
                if getattr(settings, 'RECOGNITION_ENGINE', 'yolo') == 'arcface':
                    try:
                        from app.utils.embedding_engine import compute_embedding, get_gallery
                        result = compute_embedding(pi.image.path)
                        if result is not None:
                            emb, det_conf, _ = result
                            PersonEmbedding.objects.create(
                                person=person,
                                source_image=pi,
                                embedding=emb.tolist(),
                            )
                            get_gallery().add_embedding(person.id, person.name, emb)
                            embedding_count += 1
                    except Exception as e:
                        print(f"[DatasetUpload] Auto-embedding error for {pi.image.name}: {e}")

        serializer = PersonSerializer(person, context={'request': request})
        res_status = status.HTTP_201_CREATED if created else status.HTTP_200_OK
        response_data = serializer.data
        if embedding_count > 0:
            response_data['embeddings_computed'] = embedding_count
        return Response(response_data, status=res_status)


class ClassifyUnknownsAPIView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        try:
            from app.utils.face_classifier import cluster_faces
            groups = cluster_faces()
            return Response({'status': 'success', 'groups': groups}, status=status.HTTP_200_OK)
        except Exception as e:
            return Response({'status': 'error', 'message': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class AssignClassifiedGroupAPIView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        try:
            from app.utils.face_classifier import assign_person
            group_images = request.data.get('group_images', [])
            person_name = request.data.get('person_name', '')

            if not group_images or not person_name:
                return Response({'status': 'error', 'message': 'Missing group images or person name.'}, status=status.HTTP_400_BAD_REQUEST)

            success, msg = assign_person(group_images, person_name)
            if success:
                return Response({'status': 'success', 'message': msg}, status=status.HTTP_200_OK)
            else:
                return Response({'status': 'error', 'message': msg}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as e:
            return Response({'status': 'error', 'message': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


# ==============================================================================
# REST FRAMEWORK RECOGNITION LOGS & UNKNOWN CAPTURES
# ==============================================================================

class RecognitionLogViewSet(viewsets.ReadOnlyModelViewSet):
    serializer_class = RecognitionLogSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        queryset = RecognitionLog.objects.all().order_by('-timestamp')
        status_param = self.request.query_params.get('status')
        if status_param:
            queryset = queryset.filter(status=status_param.upper())
        search = self.request.query_params.get('search')
        if search:
            queryset = queryset.filter(person_name__icontains=search)
        return queryset


class DismissUnknownLogView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def delete(self, request, log_id):
        log = get_object_or_404(RecognitionLog, id=log_id)
        if log.image_path:
            full_path = os.path.join(settings.MEDIA_ROOT, log.image_path)
            if os.path.isfile(full_path):
                try:
                    os.remove(full_path)
                except Exception as e:
                    print(f"Error deleting unknown image file: {e}")
        log.delete()
        return Response({'status': 'success', 'message': 'Log dismissed'}, status=status.HTTP_200_OK)


# ==============================================================================
# REST FRAMEWORK REPORTS & EXPORTS
# ==============================================================================


# ==============================================================================
# REST FRAMEWORK STREAMING & TELEMETRY API ENDPOINTS
# ==============================================================================

@api_view(['GET'])
@permission_classes([permissions.AllowAny])
def video_feed(request):
    src = request.GET.get('src', '0')
    model = request.GET.get('model', 'yolov8n')
    orient = request.GET.get('orient', 'normal')
    stats_key = request.GET.get('stats_key', src)
    camera_name = request.GET.get('camera_name', 'Default Camera')

    if src and src.isdigit():
        try:
            cam = Camera.objects.get(id=int(src))
            src = cam.source
            orient = cam.orientation
            camera_name = cam.name
        except Camera.DoesNotExist:
            pass

    print(f"\n[AI Stream Engine] >>> Starting Video Feed with Model: '{model}' (Camera: '{camera_name}', Src: '{src}') <<<")

    return StreamingHttpResponse(
        gen_frames(src, model, orient, stats_key, camera_name),
        content_type='multipart/x-mixed-replace; boundary=frame'
    )


@api_view(['GET'])
@permission_classes([permissions.AllowAny])
def video_stats(request):
    src = request.GET.get('src', '0')
    stats = _feed_stats.get(src, {"fps": 0, "persons": 0, "faces": 0, "names": []})
    desired = stats.get('desired_state', '')
    if desired == 'STOP' or not stats.get('is_active', True):
        return Response({
            "fps": 0.0,
            "persons": 0,
            "faces": 0,
            "names": [],
            "resolution": "0x0",
            "is_active": False,
            "desired_state": "STOP"
        })
    fps_val = stats.get('fps', 0)
    is_active = (fps_val > 0 or desired == 'START') and desired != 'STOP'
    res_data = dict(stats)
    res_data['is_active'] = is_active
    return Response(res_data)


@api_view(['POST'])
@permission_classes([permissions.AllowAny])
def start_video_feed(request):
    src = str(request.data.get('src', '0'))
    model = request.data.get('model', 'yolov8n_onnx')
    if src not in _feed_stats:
        _feed_stats[src] = {"fps": 0, "persons": 0, "faces": 0, "names": []}
    _feed_stats[src]['desired_state'] = 'START'
    _feed_stats[src]['is_active'] = True
    _feed_stats[src]['current_model'] = model
    print(f"\n[AI Stream Engine] >>> Model Selection Switched To: '{model}' for Source: '{src}' <<<")
    return Response({'status': 'started', 'src': src, 'model': model})


@api_view(['POST'])
@permission_classes([permissions.AllowAny])
def set_model(request):
    src = str(request.data.get('src', '0'))
    model = request.data.get('model', 'yolov8n_onnx')
    print(f"\n[AI Stream Engine] >>> Preferred Model Switched To: '{model}' (Stream Stopped) <<<")
    return Response({'status': 'model_updated', 'model': model})


@api_view(['POST'])
@permission_classes([permissions.AllowAny])
def stop_video_feed(request):
    src = str(request.data.get('src', '0'))
    _feed_stats[src] = {
        "fps": 0.0,
        "persons": 0,
        "faces": 0,
        "names": [],
        "resolution": "0x0",
        "is_active": False,
        "desired_state": "STOP"
    }
    if src != '0':
        _feed_stats['0'] = {
            "fps": 0.0,
            "persons": 0,
            "faces": 0,
            "names": [],
            "resolution": "0x0",
            "is_active": False,
            "desired_state": "STOP"
        }
    try:
        from app.utils.video_processor import _CAMERA_POOL
        for key in list(_CAMERA_POOL.keys()):
            cam = _CAMERA_POOL.pop(key, None)
            if cam:
                try:
                    cam.release()
                except Exception:
                    pass
    except Exception:
        pass
    return Response({'status': 'stopped', 'src': src})


@api_view(['POST'])
@permission_classes([permissions.AllowAny])
def process_client_frame(request):
    try:
        data = request.data
        img_data = data.get('image', '')
        model_name = data.get('model', 'yolov8n')
        orientation = data.get('orientation', 'normal')

        if not img_data:
            return Response({'status': 'error', 'message': 'No image data'}, status=status.HTTP_400_BAD_REQUEST)

        if ',' in img_data:
            img_data = img_data.split(',')[1]

        img_bytes = base64.b64decode(img_data)
        nparr = np.frombuffer(img_bytes, np.uint8)
        frame = cv.imdecode(nparr, cv.IMREAD_COLOR)

        if frame is None:
            return Response({'status': 'error', 'message': 'Invalid image format'}, status=status.HTTP_400_BAD_REQUEST)

        frame = rotate_and_flip_frame(frame, orientation)
        model = get_yolo_model(model_name)
        session_key = getattr(request.session, 'session_key', None) or 'client_stream'

        annotated_frame = frame.copy()
        person_count = 0
        max_confidence = 0.0

        results = model(frame, verbose=False, conf=0.45)
        for r in results:
            for box in r.boxes:
                cls_id = int(box.cls[0])
                label = model.names[cls_id]
                if label.lower() == 'person':
                    person_count += 1
                    conf = float(box.conf[0])
                    if conf > max_confidence:
                        max_confidence = conf

                    x1, y1, x2, y2 = map(int, box.xyxy[0])
                    cv.rectangle(annotated_frame, (x1, y1), (x2, y2), (0, 255, 0), 2)
                    cv.putText(annotated_frame, f"Person {conf:.2f}", (x1, max(y1-10, 20)),
                               cv.FONT_HERSHEY_SIMPLEX, 0.5, (0, 255, 0), 2)

        stats_key = f"client_{session_key}"
        _feed_stats[stats_key] = {
            "fps": 0,
            "persons": person_count
        }

        _, buffer = cv.imencode('.jpg', annotated_frame)
        processed_base64 = base64.b64encode(buffer).decode('utf-8')

        return Response({
            'status': 'success',
            'image': f"data:image/jpeg;base64,{processed_base64}",
            'faces': person_count,
            'stats_key': stats_key
        })
    except Exception as e:
        return Response({'status': 'error', 'message': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


@api_view(['GET'])
@permission_classes([permissions.AllowAny])
def test_alert(request):
    h, w = 480, 640
    frame = np.zeros((h, w, 3), dtype=np.uint8)
    cv.rectangle(frame, (100, 100), (540, 380), (0, 0, 255), 2)
    cv.putText(frame, "TEST STRANGER", (120, 90), cv.FONT_HERSHEY_SIMPLEX, 0.7, (0, 0, 255), 2, cv.LINE_AA)
    ret, buffer = cv.imencode('.jpg', frame)
    if ret:
        send_alerts(buffer.tobytes(), 'Unknown Test Target', False, 1, 0.54, camera_name="Test Camera")
        return Response({'message': 'Test alert sent successfully.'})
    return Response({'error': 'Failed to encode frame'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


@api_view(['POST'])
@permission_classes([permissions.AllowAny])
def face_verify_frame(request):
    """
    Receive a base64-encoded camera frame from the mobile app,
    run face recognition (ArcFace or YOLO), and return JWT tokens if an admin user is recognized.
    """
    try:
        img_data = request.data.get('image', '')
        if not img_data:
            return Response({'status': 'error', 'message': 'No image data provided'}, status=status.HTTP_400_BAD_REQUEST)

        # Strip data URL prefix if present
        if ',' in img_data:
            img_data = img_data.split(',')[1]

        img_bytes = base64.b64decode(img_data)
        nparr = np.frombuffer(img_bytes, np.uint8)
        frame = cv.imdecode(nparr, cv.IMREAD_COLOR)

        if frame is None:
            return Response({'status': 'error', 'message': 'Invalid image format'}, status=status.HTTP_400_BAD_REQUEST)

        use_arcface = getattr(settings, 'RECOGNITION_ENGINE', 'yolo') == 'arcface'

        if use_arcface:
            # ── ArcFace Pipeline ──
            from app.utils.embedding_engine import detect_and_recognize
            login_threshold = getattr(settings, 'ARCFACE_LOGIN_THRESHOLD', 0.60)
            detections = detect_and_recognize(frame, threshold=login_threshold, max_faces=1)

            if not detections:
                return Response({
                    'status': 'no_face',
                    'message': 'No face detected in frame',
                    'faces': 0
                })

            best = detections[0]
            recognized_name = best['person_name']
            best_conf = best['recognition_similarity']
            det_conf = best['detection_confidence']

            if not recognized_name:
                return Response({
                    'status': 'unrecognized',
                    'message': f'Face detected but no matching identity found (similarity: {round(best_conf * 100, 1)}%)',
                    'faces': 1,
                    'detection_confidence': round(det_conf, 3),
                    'recognition_similarity': round(best_conf, 3)
                })
        else:
            # ── YOLO Fallback ──
            model = get_yolo_model('yolov8n')
            results = model(frame, conf=0.25, verbose=False)

            best_conf = 0.0
            recognized_name = None

            for box in results[0].boxes:
                cls_id = int(box.cls[0])
                conf = float(box.conf[0])
                if conf > best_conf:
                    best_conf = conf
                    recognized_name = model.names[cls_id]

            if not recognized_name:
                return Response({
                    'status': 'no_face',
                    'message': 'No face detected in frame',
                    'faces': 0
                })

        # Check if recognized person is a registered admin user
        try:
            user = User.objects.get(username__iexact=recognized_name)
            is_admin = user.is_staff or user.is_superuser
        except User.DoesNotExist:
            return Response({
                'status': 'unrecognized',
                'message': f'Detected "{recognized_name}" but no matching user found',
                'faces': 1,
                'confidence': round(best_conf, 3)
            })

        if not is_admin:
            return Response({
                'status': 'blocked',
                'message': f'User "{recognized_name}" is not authorized for face login',
                'faces': 1,
                'confidence': round(best_conf, 3)
            })

        min_login_conf = getattr(settings, 'ARCFACE_LOGIN_THRESHOLD', 0.60) if use_arcface else 0.50
        if best_conf < min_login_conf:
            return Response({
                'status': 'low_confidence',
                'message': f'Detected "{recognized_name}" but confidence too low ({round(best_conf * 100, 1)}%)',
                'faces': 1,
                'confidence': round(best_conf, 3),
                'username': recognized_name
            })

        # Success — generate JWT tokens for the recognized admin
        from rest_framework_simplejwt.tokens import RefreshToken
        from .serializers import UserSerializer
        refresh = RefreshToken.for_user(user)

        return Response({
            'status': 'success',
            'username': user.username,
            'faces': 1,
            'confidence': round(best_conf, 3),
            'access': str(refresh.access_token),
            'refresh': str(refresh),
            'user': UserSerializer(user).data
        })

    except Exception as e:
        return Response({'status': 'error', 'message': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


@api_view(['GET'])
@permission_classes([permissions.AllowAny])
def face_login_feed(request):
    token = request.GET.get('token')
    return StreamingHttpResponse(
        gen_face_login_frames(token),
        content_type='multipart/x-mixed-replace; boundary=frame'
    )


@api_view(['GET'])
@permission_classes([permissions.AllowAny])
def face_login_check(request):
    token = request.GET.get('token')
    if not token:
        return Response({'status': 'pending'})

    # Check Django Cache first (Fix 7), fallback to in-memory dict
    cache_key = f'face_login_{token}'
    username = cache.get(cache_key) or _face_login_verified.get(token)
    if username:
        cache.delete(cache_key)
        _face_login_verified.pop(token, None)
        try:
            user = User.objects.get(username=username)
            from rest_framework_simplejwt.tokens import RefreshToken
            refresh = RefreshToken.for_user(user)
            return Response({
                'status': 'success',
                'username': username,
                'access': str(refresh.access_token),
                'refresh': str(refresh),
                'user': UserSerializer(user).data
            })
        except User.DoesNotExist:
            return Response({'status': 'error', 'message': 'User not found'}, status=status.HTTP_404_NOT_FOUND)
    return Response({'status': 'pending'})


# ==============================================================================
# REST FRAMEWORK REPORTS EXPORT & STATS API VIEWS
# ==============================================================================

class ReportsExportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        search_query = request.GET.get('search', '').strip()
        status_query = request.GET.get('status', '')
        time_range = request.GET.get('time_range', 'all')
        person_filter = request.GET.get('person_name', '').strip() or request.GET.get('export_person', '').strip()

        export_date = request.GET.get('export_date', '').strip()
        start_date_str = request.GET.get('start_date', '').strip()
        end_date_str = request.GET.get('end_date', '').strip()
        export_time_start = request.GET.get('export_time_start', '').strip()
        export_time_end = request.GET.get('export_time_end', '').strip()

        logs = RecognitionLog.objects.all().order_by('-timestamp')

        if search_query:
            logs = logs.filter(person_name__icontains=search_query)
        if status_query:
            logs = logs.filter(status=status_query)

        now = timezone.now()

        if export_date:
            logs = logs.filter(timestamp__date=export_date)
            title_suffix = f"Daily Report ({export_date})"
        elif start_date_str and end_date_str:
            try:
                start_d = datetime.strptime(start_date_str, '%Y-%m-%d').date()
                end_d = datetime.strptime(end_date_str, '%Y-%m-%d').date()
                logs = logs.filter(timestamp__date__range=[start_d, end_d])
                title_suffix = f"Custom Report ({start_date_str} to {end_date_str})"
            except ValueError:
                title_suffix = "Custom Report"
        elif start_date_str:
            try:
                start_d = datetime.strptime(start_date_str, '%Y-%m-%d').date()
                logs = logs.filter(timestamp__date__gte=start_d)
                title_suffix = f"Report From {start_date_str}"
            except ValueError:
                title_suffix = "Report"
        elif end_date_str:
            try:
                end_d = datetime.strptime(end_date_str, '%Y-%m-%d').date()
                logs = logs.filter(timestamp__date__lte=end_d)
                title_suffix = f"Report Until {end_date_str}"
            except ValueError:
                title_suffix = "Report"
        else:
            if time_range in ['daily', 'today']:
                logs = logs.filter(timestamp__date=now.date())
                title_suffix = f"Daily Report ({now.strftime('%d-%m-%Y')})"
            elif time_range == 'weekly':
                seven_days_ago = now - timedelta(days=7)
                logs = logs.filter(timestamp__gte=seven_days_ago)
                title_suffix = "Weekly Report (Last 7 Days)"
            elif time_range == 'monthly':
                logs = logs.filter(timestamp__year=now.year, timestamp__month=now.month)
                title_suffix = f"Monthly Report ({now.strftime('%B %Y')})"
            elif time_range == 'yearly':
                logs = logs.filter(timestamp__year=now.year)
                title_suffix = f"Yearly Report ({now.year})"
            else:
                title_suffix = "All Historical Detection Records"

        # Handle time of day range
        if export_time_start:
            try:
                t_start = datetime.strptime(export_time_start, '%H:%M').time()
                logs = logs.filter(timestamp__time__gte=t_start)
            except ValueError:
                pass
        if export_time_end:
            try:
                t_end = datetime.strptime(export_time_end, '%H:%M').time()
                logs = logs.filter(timestamp__time__lte=t_end)
            except ValueError:
                pass

        # Handle person filter
        if person_filter and person_filter.lower() != 'all':
            if person_filter.lower() == 'unknown':
                logs = logs.filter(person_name__in=[None, '', 'Unknown'])
            else:
                logs = logs.filter(person_name__iexact=person_filter)

        mapped_records = []
        for log in logs:
            mapped_records.append({
                'date': log.timestamp.date() if log.timestamp else '',
                'camera_name': log.camera_name if log.camera_name else 'Default Camera',
                'person_name': log.person_name if log.person_name else ('Unknown Target' if log.status == 'UNKNOWN' else 'Unknown'),
                'status': log.status,
                'entry_time': log.timestamp,
                'exit_time': log.timestamp,
                'frequency': 1,
                'max_confidence': float(log.confidence) if log.confidence else 0.0
            })

        wb = Workbook()

        # Sheet 1: Summary Dashboard
        ws_summary = wb.active
        ws_summary.title = "Summary"
        _fill_summary_sheet(ws_summary, mapped_records, title_suffix)

        # Sheet 2: Detection Records
        ws_records = wb.create_sheet(title="Detection Records")
        title_text = f"SMART SIGHT — {title_suffix.upper()}"
        _fill_excel_worksheet(ws_records, mapped_records, title_text)

        # Sheet 3: Camera Analytics (if records exist)
        if mapped_records:
            ws_cameras = wb.create_sheet(title="Camera Analytics")
            _fill_camera_analytics_sheet(ws_cameras, mapped_records)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"SmartSight_{time_range}_Report_{now.strftime('%Y%m%d_%H%M%S')}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        response['Access-Control-Expose-Headers'] = 'Content-Disposition'

        wb.save(response)
        return response


class ReportsStatsAPIView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        logs = RecognitionLog.objects.all()
        total_count = logs.count()
        known_count = logs.filter(status='KNOWN').count()
        unknown_count = total_count - known_count

        return Response({
            'total_detections': total_count,
            'known_detections': known_count,
            'unknown_detections': unknown_count,
        })


# Notification API Endpoints
class NotificationListAPIView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        now = timezone.now()
        # Auto-expire any pending notifications older than 15 minutes
        timeout_threshold = now - timezone.timedelta(minutes=15)
        stale_pending = Notification.objects.filter(status='PENDING', created_at__lt=timeout_threshold)
        if stale_pending.exists():
            stale_pending.update(status='EXPIRED', is_read=True, processed_at=now)

        # Ensure all already expired, approved, or cancelled alerts are marked read
        Notification.objects.filter(status__in=['EXPIRED', 'APPROVED', 'CANCELLED'], is_read=False).update(is_read=True)

        notifications = Notification.objects.all()[:50]
        serializer = NotificationSerializer(notifications, many=True)
        # Unread count strictly counts only active, unhandled PENDING alerts
        unread_count = Notification.objects.filter(status='PENDING', is_read=False).count()
        return Response({
            'notifications': serializer.data,
            'unread_count': unread_count
        })

    def delete(self, request):
        Notification.objects.all().delete()
        return Response({'status': 'all_deleted', 'message': 'All notifications deleted successfully.'})


class NotificationActionAPIView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def delete(self, request, pk):
        try:
            notification = Notification.objects.get(pk=pk)
            notification.delete()
            return Response({'status': 'deleted', 'id': pk})
        except Notification.DoesNotExist:
            return Response({'error': 'Notification not found'}, status=status.HTTP_404_NOT_FOUND)

    def post(self, request, pk):
        try:
            notification = Notification.objects.get(pk=pk)
        except Notification.DoesNotExist:
            return Response({'error': 'Notification not found'}, status=status.HTTP_404_NOT_FOUND)

        action = request.data.get('action')  # 'approve' or 'cancel'
        source = request.data.get('source', 'APP')  # 'APP' or 'TELEGRAM'
        now = timezone.now()

        # Anti-Clash & Concurrent Priority Logic:
        # If notification was already processed:
        if notification.status != 'PENDING':
            # Calculate time difference since first action
            time_diff = (now - notification.processed_at).total_seconds() if notification.processed_at else 999
            
            # Concurrent Conflict Rule:
            # If Telegram processed it recently (within 15 seconds) and Mobile App responds simultaneously,
            # Mobile App gets top priority and OVERRIDES Telegram!
            if source == 'APP' and notification.action_source == 'TELEGRAM' and time_diff <= 15:
                pass  # Allow Mobile App override
            else:
                # Normal operation: Preserve the original action source (APP or TELEGRAM)
                return Response({
                    'status': 'already_processed',
                    'message': f'Notification was already processed via {notification.get_action_source_display() if hasattr(notification, "get_action_source_display") else notification.action_source}',
                    'notification_status': notification.status,
                    'action_source': notification.action_source
                }, status=status.HTTP_409_CONFLICT)

        # Apply action with source tracking
        new_status = 'APPROVED' if action == 'approve' else ('CANCELLED' if action == 'cancel' else None)
        if new_status:
            if action == 'approve':
                # Use structured fields directly (Fix 8), with regex fallback for legacy records
                person_name = notification.alert_person_name or "Unknown"
                camera_name = notification.alert_camera_name or "Default Camera"
                confidence = notification.alert_confidence if notification.alert_confidence is not None else 0.0

                if not notification.alert_person_name and not notification.alert_camera_name:
                    import re
                    pm = re.search(r"\((.*?)\)", notification.title)
                    if pm:
                        person_name = pm.group(1).strip()
                        if 'unknown' in person_name.lower():
                            person_name = 'Unknown'
                    cm = re.search(r"on (.*?) \(Confidence: (.*?)%\)", notification.message)
                    if cm:
                        camera_name = cm.group(1)
                        try:
                            confidence = float(cm.group(2)) / 100.0
                        except Exception:
                            pass

                relative_img = notification.image_url.replace('/media/', '') if notification.image_url else None

                try:
                    RecognitionLog.objects.create(
                        person_name=person_name,
                        camera_name=camera_name,
                        confidence=confidence,
                        status='UNKNOWN',
                        image_path=relative_img
                    )
                    print(f"[Mobile App] Approved and logged: {person_name}")
                except Exception as e:
                    print(f"[Mobile App] Error creating RecognitionLog: {e}")

            notification.status = new_status
            notification.action_source = source
            notification.processed_at = now
            notification.is_read = True
            notification.save()
            return Response({
                'status': 'success',
                'notification_status': new_status,
                'action_source': source
            })

        return Response({'error': 'Invalid action'}, status=status.HTTP_400_BAD_REQUEST)


class RegisterPushTokenAPIView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        token = request.data.get('device_token')
        platform = request.data.get('platform', 'android')
        if not token:
            return Response({'error': 'device_token is required'}, status=status.HTTP_400_BAD_REQUEST)
        
        device_token, created = DevicePushToken.objects.get_or_create(
            token=token,
            defaults={'platform': platform, 'user': request.user if request.user.is_authenticated else None}
        )
        return Response({
            'status': 'success',
            'message': 'FCM Push Token registered successfully',
            'created': created
        })


@api_view(['POST'])
@permission_classes([permissions.AllowAny])
def webrtc_offer(request):
    import asyncio
    from app.utils.webrtc_processor import process_webrtc_offer

    sdp_offer = request.data.get('sdp')
    camera_src = request.data.get('src', '0')
    model_name = request.data.get('model', 'yolov8n_onnx')

    if not sdp_offer:
        return Response({'error': 'SDP offer is required'}, status=status.HTTP_400_BAD_REQUEST)

    try:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        res = loop.run_until_complete(process_webrtc_offer(sdp_offer, camera_src, model_name))
        loop.close()
        return Response(res, status=status.HTTP_200_OK)
    except Exception as e:
        return Response({'status': 'fallback', 'message': str(e)}, status=status.HTTP_200_OK)


# ==============================================================================
# GALLERY MANAGEMENT API VIEWS (ArcFace Embedding Engine)
# ==============================================================================

class GalleryStatusAPIView(APIView):
    """Returns the current state of the ArcFace embedding gallery."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        engine = getattr(settings, 'RECOGNITION_ENGINE', 'yolo')
        if engine != 'arcface':
            return Response({
                'engine': engine,
                'status': 'inactive',
                'message': 'ArcFace engine is not active. Set RECOGNITION_ENGINE=arcface in settings.'
            })

        from app.utils.embedding_engine import get_gallery
        gallery = get_gallery()
        stats = gallery.stats()
        return Response({
            'engine': engine,
            'status': 'active',
            'threshold': getattr(settings, 'ARCFACE_SIMILARITY_THRESHOLD', 0.55),
            'login_threshold': getattr(settings, 'ARCFACE_LOGIN_THRESHOLD', 0.60),
            'detection_backend': getattr(settings, 'ARCFACE_DETECTION_BACKEND', 'retinaface'),
            **stats,
        })


class GalleryRebuildAPIView(APIView):
    """Force-reload the in-memory embedding gallery from the database."""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        engine = getattr(settings, 'RECOGNITION_ENGINE', 'yolo')
        if engine != 'arcface':
            return Response({'error': 'ArcFace engine is not active.'}, status=status.HTTP_400_BAD_REQUEST)

        from app.utils.embedding_engine import get_gallery
        gallery = get_gallery()
        gallery.reload()
        stats = gallery.stats()
        return Response({
            'status': 'success',
            'message': 'Gallery rebuilt from database.',
            **stats,
        })


class ComputePersonEmbeddingsAPIView(APIView):
    """Compute (or recompute) ArcFace embeddings for all images of a person."""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, person_id):
        engine = getattr(settings, 'RECOGNITION_ENGINE', 'yolo')
        if engine != 'arcface':
            return Response({'error': 'ArcFace engine is not active.'}, status=status.HTTP_400_BAD_REQUEST)

        from app.utils.embedding_engine import compute_person_embeddings
        try:
            person = Person.objects.get(id=person_id)
        except Person.DoesNotExist:
            return Response({'error': 'Person not found.'}, status=status.HTTP_404_NOT_FOUND)

        recompute = request.data.get('recompute', False)
        if recompute:
            # Delete existing embeddings and recompute all
            PersonEmbedding.objects.filter(person=person).delete()

        computed, skipped, errors = compute_person_embeddings(person_id)
        return Response({
            'status': 'success',
            'person_name': person.name,
            'computed': computed,
            'skipped': skipped,
            'errors': errors,
        })


@api_view(['GET', 'HEAD', 'OPTIONS'])
@permission_classes([permissions.AllowAny])
def health_check(request):
    """
    Lightweight health check endpoint for desktop/mobile app connectivity tests.
    """
    return Response({
        'status': 'ok',
        'app': 'SmartSight',
        'timestamp': timezone.now().isoformat()
    })

