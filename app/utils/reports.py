from collections import Counter
from django.utils import timezone
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter


# ── Color Palette ──
BRAND_BLUE      = "0D6EFD"
DARK_HEADER     = "0F172A"
SUBHEADER_BG    = "1E293B"
ROW_EVEN        = "F8FAFC"
ROW_ODD         = "FFFFFF"
KNOWN_GREEN     = "15803D"
UNKNOWN_RED     = "B91C1C"
ACCENT_EMERALD  = "10B981"
BORDER_SUBTLE   = "CBD5E1"
DARK_BG         = "1E293B"
LIGHT_BG        = "F1F5F9"
FOOTER_BG       = "0F172A"

THIN_BORDER = Border(
    left=Side(style='thin', color=BORDER_SUBTLE),
    right=Side(style='thin', color=BORDER_SUBTLE),
    top=Side(style='thin', color=BORDER_SUBTLE),
    bottom=Side(style='thin', color=BORDER_SUBTLE)
)


def _fill_summary_sheet(ws, reports, title_suffix):
    """Create a branded summary dashboard sheet."""
    now = timezone.localtime(timezone.now())
    total = len(reports)
    known = sum(1 for r in reports if r.get('status') == 'KNOWN')
    unknown = total - known
    known_pct = round((known / total) * 100) if total > 0 else 0
    unknown_pct = 100 - known_pct if total > 0 else 0

    ws.sheet_properties.tabColor = BRAND_BLUE

    # ── Title Banner ──
    ws.merge_cells('A1:F1')
    ws['A1'] = f"SMART SIGHT — {title_suffix.upper()}"
    ws['A1'].font = Font(name='Inter', size=18, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color=BRAND_BLUE, end_color=BRAND_BLUE, fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 50

    # ── Subtitle ──
    ws.merge_cells('A2:F2')
    ws['A2'] = f"Generated: {now.strftime('%A, %d %B %Y  •  %I:%M %p')}"
    ws['A2'].font = Font(name='Inter', size=10, italic=True, color='64748B')
    ws['A2'].fill = PatternFill(start_color=LIGHT_BG, end_color=LIGHT_BG, fill_type='solid')
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 28

    # Spacer
    ws.row_dimensions[3].height = 10

    # ── Quick Stats Section ──
    ws.merge_cells('A4:F4')
    ws['A4'] = "SUMMARY STATISTICS"
    ws['A4'].font = Font(name='Inter', size=12, bold=True, color='FFFFFF')
    ws['A4'].fill = PatternFill(start_color=DARK_HEADER, end_color=DARK_HEADER, fill_type='solid')
    ws['A4'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[4].height = 32

    stats = [
        ('Total Detections', str(total), BRAND_BLUE),
        ('Known Persons', f"{known}  ({known_pct}%)", KNOWN_GREEN),
        ('Unknown Persons', f"{unknown}  ({unknown_pct}%)", UNKNOWN_RED),
    ]

    for i, (label, value, color) in enumerate(stats):
        row = 5 + i
        ws.cell(row=row, column=1, value='').border = THIN_BORDER
        ws.cell(row=row, column=2, value=label)
        ws.cell(row=row, column=2).font = Font(name='Inter', size=11, bold=True, color='334155')
        ws.cell(row=row, column=2).alignment = Alignment(horizontal='left', vertical='center')
        ws.cell(row=row, column=2).border = THIN_BORDER

        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        ws.cell(row=row, column=3, value=value)
        ws.cell(row=row, column=3).font = Font(name='Inter', size=12, bold=True, color=color)
        ws.cell(row=row, column=3).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=row, column=3).border = THIN_BORDER

        fill_color = ROW_EVEN if i % 2 == 0 else ROW_ODD
        for c in range(1, 7):
            ws.cell(row=row, column=c).fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
            ws.cell(row=row, column=c).border = THIN_BORDER
        ws.row_dimensions[row].height = 28

    # Spacer
    ws.row_dimensions[8].height = 10

    # ── Top Detected Persons ──
    person_counts = Counter()
    person_status = {}
    for r in reports:
        name = r.get('person_name') or 'Unknown'
        person_counts[name] += 1
        person_status[name] = r.get('status', 'UNKNOWN')

    top_persons = person_counts.most_common(10)

    if top_persons:
        ws.merge_cells('A9:F9')
        ws['A9'] = "TOP DETECTED PERSONS"
        ws['A9'].font = Font(name='Inter', size=12, bold=True, color='FFFFFF')
        ws['A9'].fill = PatternFill(start_color=DARK_HEADER, end_color=DARK_HEADER, fill_type='solid')
        ws['A9'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[9].height = 32

        tp_headers = ['#', 'Person Name', 'Detections', 'Status', '', '']
        for ci, h in enumerate(tp_headers, 1):
            cell = ws.cell(row=10, column=ci, value=h)
            cell.font = Font(name='Inter', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color=SUBHEADER_BG, end_color=SUBHEADER_BG, fill_type='solid')
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = THIN_BORDER
        ws.row_dimensions[10].height = 24

        for idx, (name, count) in enumerate(top_persons):
            row = 11 + idx
            st = person_status.get(name, 'UNKNOWN')
            status_text = "Known" if st == "KNOWN" else "Unknown"
            status_color = KNOWN_GREEN if st == "KNOWN" else UNKNOWN_RED

            ws.cell(row=row, column=1, value=idx + 1)
            ws.cell(row=row, column=2, value=name)
            ws.cell(row=row, column=3, value=count)
            ws.cell(row=row, column=4, value=status_text)
            ws.cell(row=row, column=4).font = Font(name='Inter', size=10, bold=True, color=status_color)

            fill = ROW_EVEN if idx % 2 == 0 else ROW_ODD
            for c in range(1, 7):
                cell = ws.cell(row=row, column=c)
                cell.fill = PatternFill(start_color=fill, end_color=fill, fill_type='solid')
                cell.border = THIN_BORDER
                if c != 4:
                    cell.font = Font(name='Inter', size=10)
                cell.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[row].height = 22

    # ── Footer ──
    footer_row = max(21, 11 + len(top_persons) + 2)
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=6)
    ws.cell(row=footer_row, column=1, value="Powered by Smart Sight AI  •  Automated Surveillance Intelligence")
    ws.cell(row=footer_row, column=1).font = Font(name='Inter', size=9, italic=True, color='94A3B8')
    ws.cell(row=footer_row, column=1).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=footer_row, column=1).fill = PatternFill(start_color=FOOTER_BG, end_color=FOOTER_BG, fill_type='solid')
    for c in range(1, 7):
        ws.cell(row=footer_row, column=c).fill = PatternFill(start_color=FOOTER_BG, end_color=FOOTER_BG, fill_type='solid')
    ws.row_dimensions[footer_row].height = 30

    # Column widths
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12


def _fill_excel_worksheet(ws, reports, title_text):
    """Fill a worksheet with styled detection records."""
    ws.sheet_properties.tabColor = "10B981"
    ws.views.sheetView[0].showGridLines = False

    # ── Title Banner ──
    ws.append([title_text])
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.font = Font(name="Inter", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color=BRAND_BLUE, end_color=BRAND_BLUE, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 45

    # ── Subtitle ──
    now = timezone.localtime(timezone.now())
    ws.append([f"Generated: {now.strftime('%d %b %Y, %I:%M %p')}  •  {len(reports)} Records"])
    ws.merge_cells("A2:H2")
    ws["A2"].font = Font(name="Inter", size=10, italic=True, color="64748B")
    ws["A2"].fill = PatternFill(start_color=LIGHT_BG, end_color=LIGHT_BG, fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 26

    # ── Headers ──
    headers = ["#", "Date", "Camera", "Person Name", "Classification", "Entry Time", "Exit Time", "Confidence"]
    ws.append(headers)
    ws.row_dimensions[3].height = 28

    header_fill = PatternFill(start_color=DARK_HEADER, end_color=DARK_HEADER, fill_type="solid")
    header_font = Font(name="Inter", size=10, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_num in range(1, 9):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = THIN_BORDER

    # ── Data Rows ──
    row_idx = 4
    for idx, rep in enumerate(reports):
        rep_date = rep['date']
        if isinstance(rep_date, str):
            date_str = rep_date
        else:
            date_str = rep_date.strftime('%d/%m/%Y')

        name = rep['person_name'] if rep['person_name'] else "Unknown Person"
        status_disp = "Known" if rep['status'] == "KNOWN" else "Unknown"
        camera_disp = rep.get('camera_name', 'Default Camera')
        entry_local = timezone.localtime(rep['entry_time'])
        exit_local  = timezone.localtime(rep['exit_time'])
        entry_str = entry_local.strftime('%I:%M:%S %p')
        exit_str  = exit_local.strftime('%I:%M:%S %p')
        max_conf = f"{rep['max_confidence'] * 100:.1f}%" if rep.get('max_confidence') and rep['max_confidence'] <= 1.0 else f"{rep.get('max_confidence', 0):.1f}%"

        ws.append([idx + 1, date_str, camera_disp, name, status_disp, entry_str, exit_str, max_conf])
        ws.row_dimensions[row_idx].height = 22

        row_fill = PatternFill(start_color=ROW_EVEN if row_idx % 2 == 0 else ROW_ODD, fill_type="solid")
        status_color = KNOWN_GREEN if rep['status'] == "KNOWN" else UNKNOWN_RED
        status_fill = PatternFill(
            start_color="DCFCE7" if rep['status'] == "KNOWN" else "FEE2E2",
            end_color="DCFCE7" if rep['status'] == "KNOWN" else "FEE2E2",
            fill_type="solid"
        )

        for col_num in range(1, 9):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.border = THIN_BORDER

            if col_num == 5:  # Classification column
                cell.fill = status_fill
                cell.font = Font(name="Inter", size=10, bold=True, color=status_color)
            else:
                cell.fill = row_fill
                cell.font = Font(name="Inter", size=10)

            if col_num in [1, 2, 3, 5, 6, 7, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        row_idx += 1

    # ── Footer Summary Row ──
    if reports:
        known_count = sum(1 for r in reports if r['status'] == 'KNOWN')
        footer_text = f"Total: {len(reports)} records  |  Known: {known_count}  |  Unknown: {len(reports) - known_count}"
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=8)
        ws.cell(row=row_idx, column=1, value=footer_text)
        ws.cell(row=row_idx, column=1).font = Font(name='Inter', size=10, bold=True, color='FFFFFF')
        ws.cell(row=row_idx, column=1).fill = PatternFill(start_color=FOOTER_BG, end_color=FOOTER_BG, fill_type='solid')
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal='center', vertical='center')
        for c in range(1, 9):
            ws.cell(row=row_idx, column=c).fill = PatternFill(start_color=FOOTER_BG, end_color=FOOTER_BG, fill_type='solid')
            ws.cell(row=row_idx, column=c).border = THIN_BORDER
        ws.row_dimensions[row_idx].height = 28

    # ── Auto-fit Column Widths ──
    min_widths = {'A': 5, 'B': 14, 'C': 16, 'D': 22, 'E': 16, 'F': 16, 'G': 16, 'H': 14}
    for col in ws.columns:
        col_letter = None
        for cell in col:
            if hasattr(cell, 'column_letter'):
                col_letter = cell.column_letter
                break
        if not col_letter:
            continue

        max_len = 0
        for cell in col:
            if cell.row <= 2:
                continue
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, min_widths.get(col_letter, 12))

    # Freeze panes at row 4
    ws.freeze_panes = 'A4'


def _fill_camera_analytics_sheet(ws, reports):
    """Create a camera breakdown analytics sheet."""
    ws.sheet_properties.tabColor = "F59E0B"
    ws.views.sheetView[0].showGridLines = False

    camera_stats = {}
    for r in reports:
        cam = r.get('camera_name', 'Default Camera')
        if cam not in camera_stats:
            camera_stats[cam] = {'total': 0, 'known': 0, 'unknown': 0}
        camera_stats[cam]['total'] += 1
        if r['status'] == 'KNOWN':
            camera_stats[cam]['known'] += 1
        else:
            camera_stats[cam]['unknown'] += 1

    # Title
    ws.merge_cells('A1:E1')
    ws['A1'] = "SMART SIGHT — CAMERA ANALYTICS"
    ws['A1'].font = Font(name='Inter', size=14, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 40

    # Spacer
    ws.row_dimensions[2].height = 8

    # Headers
    cam_headers = ['Camera Name', 'Total Detections', 'Known', 'Unknown', 'Known %']
    for ci, h in enumerate(cam_headers, 1):
        cell = ws.cell(row=3, column=ci, value=h)
        cell.font = Font(name='Inter', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color=DARK_HEADER, end_color=DARK_HEADER, fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = THIN_BORDER
    ws.row_dimensions[3].height = 26

    # Data
    sorted_cams = sorted(camera_stats.items(), key=lambda x: x[1]['total'], reverse=True)
    for idx, (cam_name, stats) in enumerate(sorted_cams):
        row = 4 + idx
        known_pct = round((stats['known'] / stats['total']) * 100) if stats['total'] > 0 else 0

        ws.cell(row=row, column=1, value=cam_name)
        ws.cell(row=row, column=2, value=stats['total'])
        ws.cell(row=row, column=3, value=stats['known'])
        ws.cell(row=row, column=4, value=stats['unknown'])
        ws.cell(row=row, column=5, value=f"{known_pct}%")

        fill = ROW_EVEN if idx % 2 == 0 else ROW_ODD
        for c in range(1, 6):
            cell = ws.cell(row=row, column=c)
            cell.fill = PatternFill(start_color=fill, end_color=fill, fill_type='solid')
            cell.border = THIN_BORDER
            cell.font = Font(name='Inter', size=10)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[row].height = 22

    # Column widths
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12


def generate_reports_excel(reports, title_suffix="FULL REPORT"):
    """Master Excel generator function for Smart Sight analytics reports."""
    from openpyxl import Workbook
    import io

    wb = Workbook()

    # Sheet 1: Executive Summary
    ws1 = wb.active
    ws1.title = "Executive Summary"
    _fill_summary_sheet(ws1, reports, title_suffix)

    # Sheet 2: Detection Logs
    ws2 = wb.create_sheet(title="Detection Logs")
    _fill_excel_worksheet(ws2, reports, f"SMART SIGHT — {title_suffix}")

    # Sheet 3: Camera Analytics
    ws3 = wb.create_sheet(title="Camera Analytics")
    _fill_camera_analytics_sheet(ws3, reports)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
