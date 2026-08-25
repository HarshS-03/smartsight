import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_bg(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

# -------------------------------------------------------------
# 1. INSTITUTIONAL & DEPARTMENT PERMISSION LETTER (1 PAGE)
# (Signatures: Mentor, HOD, Principal)
# -------------------------------------------------------------
def build_institutional_permission_doc(filename="s:/BKP/smartsight/Institutional_Permission_Letter.docx"):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10)
    font.color.rgb = RGBColor(30, 30, 30)

    # Header
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.space_before = Pt(0)
    p_hdr.paragraph_format.space_after = Pt(2)
    r1 = p_hdr.add_run("[COLLEGE / UNIVERSITY / INSTITUTION NAME]\n")
    r1.bold = True
    r1.font.size = Pt(11.5)
    r1.font.color.rgb = RGBColor(15, 35, 70)
    r2 = p_hdr.add_run("Department of Computer Science & Engineering / Artificial Intelligence\nAcademic Project & Research Cell")
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(90, 90, 90)

    # Date
    p_dt = doc.add_paragraph()
    p_dt.paragraph_format.space_after = Pt(4)
    r_dt = p_dt.add_run("Date: _____ / _____ / 2026")
    r_dt.bold = True
    r_dt.font.size = Pt(9.5)

    # To Section
    p_to = doc.add_paragraph()
    p_to.paragraph_format.space_after = Pt(5)
    p_to.paragraph_format.line_spacing = 1.12
    p_to.add_run("To,\nThe Principal / Head of the Department (HOD),\n[Department Name / Respective Department],\n[College / Institution Name]")

    # Subject
    p_sbj = doc.add_paragraph()
    p_sbj.paragraph_format.space_before = Pt(2)
    p_sbj.paragraph_format.space_after = Pt(6)
    p_sbj.paragraph_format.line_spacing = 1.12
    r_sl = p_sbj.add_run("SUBJECT: ")
    r_sl.bold = True
    r_sl.font.color.rgb = RGBColor(15, 35, 70)
    r_st = p_sbj.add_run("Request for Permission to Collect Image Dataset across Departments for Academic Project ")
    r_sp = p_sbj.add_run("“SmartSight”")
    r_sp.bold = True

    # Body
    p_b1 = doc.add_paragraph()
    p_b1.paragraph_format.space_after = Pt(5)
    p_b1.paragraph_format.line_spacing = 1.14
    p_b1.add_run(
        "Respected Sir/Madam,\n\n"
        "We, the student researchers from the Department of Computer Science & Engineering / AI, are developing an advanced academic project titled "
    )
    r_proj = p_b1.add_run("“SmartSight: Real-Time AI Face Recognition, Detection, and Surveillance Platform”")
    r_proj.bold = True
    p_b1.add_run(
        " under faculty mentorship. To train and evaluate our Computer Vision models (YOLOv8 & DeepFace/ArcFace) on diverse real-world samples, "
        "we require permission to collect face image photographs from "
    )
    r_stu = p_b1.add_run("student volunteers (~100 students across Course, Year, and Batch)")
    r_stu.bold = True
    p_b1.add_run(" from various departments.")

    p_b2 = doc.add_paragraph()
    p_b2.paragraph_format.space_after = Pt(5)
    p_b2.paragraph_format.line_spacing = 1.14
    r_vl = p_b2.add_run("Voluntary Participation: ")
    r_vl.bold = True
    p_b2.add_run(
        "Participation is entirely voluntary. Photographs will be collected only from students who willingly provide their images for this research without any academic compulsion."
    )

    p_b3 = doc.add_paragraph()
    p_b3.paragraph_format.space_after = Pt(3)
    p_b3.paragraph_format.line_spacing = 1.14
    p_b3.add_run(
        "We formally undertake and guarantee full adherence to ethical guidelines and institutional data protection:"
    )

    points = [
        ("Strict Academic & Research Scope: ", "Images will be utilized solely for model training and accuracy testing within the SmartSight project."),
        ("No Open Public Sharing / Leakage: ", "Images will never be shared on public platforms or handed to third parties."),
        ("Research Publication & Authorized Access Only: ", "If this research is published as a research paper / academic publication, the dataset will NOT be made public to everyone; safe, encrypted access will be granted strictly and exclusively to authorized researchers upon verified request."),
        ("Strict Privacy & Zero Misuse: ", "There will be no misuse, unauthorized profiling, or unethical processing of images."),
        ("Safe Deletion Protocol: ", "Upon completion of project evaluation, raw images will be safely purged as per institutional norms.")
    ]

    for title, desc in points:
        p_pt = doc.add_paragraph(style='List Bullet')
        p_pt.paragraph_format.space_after = Pt(2)
        p_pt.paragraph_format.line_spacing = 1.1
        r_t = p_pt.add_run(title)
        r_t.bold = True
        r_d = p_pt.add_run(desc)

    p_req = doc.add_paragraph()
    p_req.paragraph_format.space_before = Pt(3)
    p_req.paragraph_format.space_after = Pt(12)
    p_req.paragraph_format.line_spacing = 1.1
    p_req.add_run("We kindly request your approval to collect the required image dataset under academic supervision.")

    # 3-Column Sign-Off: Mentor, HOD, Principal
    sig_tbl = doc.add_table(rows=1, cols=3)
    sig_tbl.autofit = False

    # 1. Mentor
    c1 = sig_tbl.cell(0, 0)
    c1.width = Inches(2.3)
    p1 = c1.paragraphs[0]
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_after = Pt(0)
    r_c1 = p1.add_run(
        "Recommended by:\n\n\n\n"
        "_____________________\n"
        "Project Mentor / Guide\n"
        "Name: _______________\n"
        "Dept.: _______________"
    )
    r_c1.font.size = Pt(9)
    r_c1.bold = True

    # 2. HOD
    c2 = sig_tbl.cell(0, 1)
    c2.width = Inches(2.3)
    p2 = c2.paragraphs[0]
    p2.paragraph_format.line_spacing = 1.15
    p2.paragraph_format.space_after = Pt(0)
    r_c2 = p2.add_run(
        "Forwarded by:\n\n\n\n"
        "_____________________\n"
        "Head of Department (HOD)\n"
        "Name: _______________\n"
        "Seal / Stamp:"
    )
    r_c2.font.size = Pt(9)
    r_c2.bold = True

    # 3. Principal
    c3 = sig_tbl.cell(0, 2)
    c3.width = Inches(2.3)
    p3 = c3.paragraphs[0]
    p3.paragraph_format.line_spacing = 1.15
    p3.paragraph_format.space_after = Pt(0)
    r_c3 = p3.add_run(
        "Approved by:\n\n\n\n"
        "_____________________\n"
        "Principal / Director\n"
        "College / Institute Seal\n"
        "Date: ____ / ____ / 2026"
    )
    r_c3.font.size = Pt(9)
    r_c3.bold = True

    doc.save(filename)
    print(f"Generated Institutional Permission Letter: {filename}")


# -------------------------------------------------------------
# 2. GENERALIZED STUDENT DATASET COLLECTION & CONSENT LETTER (1 PAGE)
# (Addressed to Class/Cohort with Course & Year at top, Lead & Mentor Signatures at bottom)
# -------------------------------------------------------------
def build_student_collection_sheet(filename="s:/BKP/smartsight/Student_Dataset_Collection_Sheet.docx"):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(9.5)
    font.color.rgb = RGBColor(30, 30, 30)

    # 1. Header (Institutional & Department)
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.space_before = Pt(0)
    p_hdr.paragraph_format.space_after = Pt(2)
    r1 = p_hdr.add_run("[COLLEGE / UNIVERSITY / INSTITUTION NAME]\n")
    r1.bold = True
    r1.font.size = Pt(11.5)
    r1.font.color.rgb = RGBColor(15, 35, 70)
    r2 = p_hdr.add_run("Department of Computer Science & Engineering / Artificial Intelligence\nAcademic Research & Project Development Cell")
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(90, 90, 90)

    # 2. Target Class / Cohort Metadata Box (Course, Year, Batch)
    meta_tbl = doc.add_table(rows=2, cols=4)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False
    set_table_borders(meta_tbl, color="B0C4DE", sz="6")

    meta_fields = [
        ("Course:", "[ B.Tech / BCA / MCA / Other ]", "Year / Semester:", "[ 1st / 2nd / 3rd / 4th Year ]"),
        ("Batch / Division:", "[ Batch A / Div 1 / 2022-26 ]", "Academic Project:", "SmartSight (AI Face Dataset)")
    ]
    col_widths = [Inches(1.5), Inches(2.2), Inches(1.5), Inches(2.0)]

    for row_idx, data in enumerate(meta_fields):
        row = meta_tbl.rows[row_idx]
        for col_idx in range(4):
            c = row.cells[col_idx]
            c.width = col_widths[col_idx]
            set_cell_margins(c, top=35, bottom=35, left=50, right=50)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if col_idx % 2 == 0:
                set_cell_bg(c, "F0F4F8")
                r = p.add_run(data[col_idx])
                r.bold = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(15, 35, 70)
            else:
                r = p.add_run(data[col_idx])
                r.font.size = Pt(8.5)

    # 3. Recipient (To)
    p_to = doc.add_paragraph()
    p_to.paragraph_format.space_before = Pt(5)
    p_to.paragraph_format.space_after = Pt(3)
    p_to.paragraph_format.line_spacing = 1.12
    r_to_title = p_to.add_run("To,\n")
    r_to_title.bold = True
    p_to.add_run("All Students / Student Volunteers,\n[Course, Year & Batch as specified above],\n[College / University / Institution Name]")

    # 4. Subject
    p_sbj = doc.add_paragraph()
    p_sbj.paragraph_format.space_before = Pt(2)
    p_sbj.paragraph_format.space_after = Pt(5)
    p_sbj.paragraph_format.line_spacing = 1.12
    r_sl = p_sbj.add_run("SUBJECT: ")
    r_sl.bold = True
    r_sl.font.color.rgb = RGBColor(15, 35, 70)
    r_st = p_sbj.add_run("Information & Voluntary Dataset Collection Notice for Academic AI Project ")
    r_sp = p_sbj.add_run("“SmartSight”")
    r_sp.bold = True

    # 5. Salutation & Introduction
    p_b1 = doc.add_paragraph()
    p_b1.paragraph_format.space_after = Pt(4)
    p_b1.paragraph_format.line_spacing = 1.14
    p_b1.add_run(
        "Dear Students / Fellow Peers,\n\n"
        "We, the student project team from the Department of Computer Science & Engineering / AI, are developing an academic research project titled "
    )
    r_pj = p_b1.add_run("“SmartSight: Real-Time AI Face Recognition, Detection, and Surveillance Platform”")
    r_pj.bold = True
    p_b1.add_run(
        " under faculty mentorship. To train, calibrate, and evaluate our AI Computer Vision models (YOLOv8 & DeepFace/ArcFace) across diverse real-world conditions, "
        "we cordially invite students from your class/batch to voluntarily participate in facial image dataset collection."
    )

    # 6. Ethical Undertakings & Privacy Notice
    p_b2 = doc.add_paragraph()
    p_b2.paragraph_format.space_after = Pt(3)
    p_b2.paragraph_format.line_spacing = 1.12
    r_u = p_b2.add_run("Student Privacy Guarantees & Ethical Guidelines:")
    r_u.bold = True
    r_u.font.color.rgb = RGBColor(15, 35, 70)

    points = [
        ("100% Voluntary Participation: ", "Participation is strictly voluntary. There is no academic compulsion, pressure, or effect on attendance/grades."),
        ("Solely Academic & Research Scope: ", "Photographs will be utilized exclusively for training and evaluating AI models within the SmartSight project."),
        ("Strict Privacy & Zero Public Sharing: ", "Your photographs will NEVER be uploaded to social media, public websites, or distributed to any third party."),
        ("Research Publication Safeguard: ", "In case this research is published in academic journals/conferences, the raw dataset will NOT be accessible to the public; only safe, encrypted access will be granted strictly to authorized researchers upon verified academic request."),
        ("Safe Deletion Protocol: ", "All raw images will be safely archived and permanently deleted upon completion of final academic project evaluation.")
    ]

    for title, desc in points:
        p_pt = doc.add_paragraph(style='List Bullet')
        p_pt.paragraph_format.space_after = Pt(2)
        p_pt.paragraph_format.line_spacing = 1.08
        r_t = p_pt.add_run(title)
        r_t.bold = True
        r_d = p_pt.add_run(desc)

    # 7. Note / Call to participate
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(3)
    p_note.paragraph_format.space_after = Pt(12)
    p_note.paragraph_format.line_spacing = 1.12
    r_nt = p_note.add_run(
        "Interested students willing to contribute their photographs for this academic research may kindly coordinate with the Project Team during the scheduled session."
    )
    r_nt.italic = True
    r_nt.font.size = Pt(9)

    # 8. Signatures Table (Lead & Mentor ONLY)
    sig_tbl = doc.add_table(rows=1, cols=2)
    sig_tbl.autofit = False

    c_left = sig_tbl.cell(0, 0)
    c_left.width = Inches(3.5)
    p_s = c_left.paragraphs[0]
    p_s.paragraph_format.line_spacing = 1.15
    p_s.paragraph_format.space_after = Pt(0)
    r_s = p_s.add_run(
        "Issued by (Project Team):\n\n\n\n"
        "_____________________________\n"
        "Project Team Lead / Coordinator\n"
        "Name: _______________________\n"
        "Roll No. / Dept: ______________\n"
        "Project: SmartSight"
    )
    r_s.font.size = Pt(9)
    r_s.bold = True

    c_right = sig_tbl.cell(0, 1)
    c_right.width = Inches(3.5)
    p_m = c_right.paragraphs[0]
    p_m.paragraph_format.line_spacing = 1.15
    p_m.paragraph_format.space_after = Pt(0)
    r_m = p_m.add_run(
        "Supervised & Approved by:\n\n\n\n"
        "_____________________________\n"
        "Project Mentor / Faculty In-Charge\n"
        "Name: _______________________\n"
        "Designation: _________________\n"
        "Department: ___________________"
    )
    r_m.font.size = Pt(9)
    r_m.bold = True

    doc.save(filename)
    print(f"Generated Generalized Student Letter: {filename}")



# -------------------------------------------------------------
# 3. MENTOR PERMISSION & UNDERTAKING LETTER (1 PAGE)
# (Signatures: Submitted by Team + Approved by Mentor)
# -------------------------------------------------------------
def build_mentor_permission_letter(filename="s:/BKP/smartsight/Dataset_Consent_and_Undertaking_Letter.docx"):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(30, 30, 30)

    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.space_before = Pt(0)
    p_header.paragraph_format.space_after = Pt(2)
    r_inst = p_header.add_run("[DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING / AI]\n[COLLEGE / UNIVERSITY / INSTITUTION NAME]")
    r_inst.bold = True
    r_inst.font.size = Pt(11.5)
    r_inst.font.color.rgb = RGBColor(15, 35, 70)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run("Academic Research & Project Development Cell")
    r_sub.font.size = Pt(9.5)
    r_sub.italic = True
    r_sub.font.color.rgb = RGBColor(100, 100, 100)

    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_after = Pt(6)
    r_dt = p_date.add_run("Date: _____ / _____ / 2026")
    r_dt.bold = True
    r_dt.font.size = Pt(10)

    p_to = doc.add_paragraph()
    p_to.paragraph_format.space_after = Pt(6)
    p_to.paragraph_format.line_spacing = 1.15
    r_to_bold = p_to.add_run("To,\n")
    r_to_bold.bold = True
    p_to.add_run("The Project Mentor / Faculty Guide,\nDepartment of [Computer Science & Engineering / AI],\n[College / University Name]")

    p_subj = doc.add_paragraph()
    p_subj.paragraph_format.space_before = Pt(3)
    p_subj.paragraph_format.space_after = Pt(8)
    p_subj.paragraph_format.line_spacing = 1.15
    r_sbj_lbl = p_subj.add_run("SUBJECT: ")
    r_sbj_lbl.bold = True
    r_sbj_txt = p_subj.add_run("Application seeking permission for Face Dataset Collection & Confidentiality Undertaking for Project ")
    r_sbj_p = p_subj.add_run("“SmartSight”")
    r_sbj_p.bold = True

    p_b1 = doc.add_paragraph()
    p_b1.paragraph_format.space_after = Pt(5)
    p_b1.paragraph_format.line_spacing = 1.15
    p_b1.add_run(
        "Respected Sir/Madam,\n\n"
        "We are currently developing our academic research project titled "
    )
    r_pj = p_b1.add_run("“SmartSight: Real-Time AI Face Recognition, Detection, and Surveillance Platform”")
    r_pj.bold = True
    p_b1.add_run(
        " under your esteemed mentorship and guidance. In order to effectively train, calibrate, and evaluate our AI models "
        "(YOLOv8 & DeepFace/ArcFace), we require a diverse dataset comprising facial photographs of "
    )
    r_vol = p_b1.add_run("approximately 100 student participants / volunteers")
    r_vol.bold = True
    p_b1.add_run(" across Course, Year, and Batch from our institution.")

    p_b2 = doc.add_paragraph()
    p_b2.paragraph_format.space_after = Pt(5)
    p_b2.paragraph_format.line_spacing = 1.15
    r_w_lbl = p_b2.add_run("Voluntary Consent: ")
    r_w_lbl.bold = True
    p_b2.add_run(
        "All images will be collected purely on a willing and voluntary basis with informed consent, without any academic compulsion, coercion, or commercial motive."
    )

    p_b3 = doc.add_paragraph()
    p_b3.paragraph_format.space_after = Pt(3)
    p_b3.paragraph_format.line_spacing = 1.15
    p_b3.add_run(
        "On behalf of the project team, we hereby formally undertake and assure that:"
    )

    points = [
        ("Strict Academic Scope: ", "Images will be used solely for model training, feature extraction, and accuracy testing within SmartSight."),
        ("No Open Public Sharing: ", "These images will NOT be shared, published, or transferred to public platforms or third parties."),
        ("Research Publication & Authorized Access Only: ", "In the event that this research work is published in an academic journal or conference, the dataset will NOT be made public for general access. Only safe, encrypted, and strictly controlled access will be provided to verified and authorized researchers upon formal academic request."),
        ("Strict Privacy & Zero Misuse: ", "There will be no misuse, unauthorized profiling, or unethical alteration of the images. Complete privacy will be strictly maintained."),
        ("Safe Deletion Protocol: ", "Upon completion of final evaluation and viva voce, raw images will be safely purged as per institutional norms.")
    ]

    for b_title, b_desc in points:
        p_pt = doc.add_paragraph(style='List Bullet')
        p_pt.paragraph_format.space_after = Pt(2.5)
        p_pt.paragraph_format.line_spacing = 1.12
        r_bt = p_pt.add_run(b_title)
        r_bt.bold = True
        r_bd = p_pt.add_run(b_desc)

    p_close = doc.add_paragraph()
    p_close.paragraph_format.space_before = Pt(3)
    p_close.paragraph_format.space_after = Pt(12)
    p_close.paragraph_format.line_spacing = 1.12
    p_close.add_run(
        "We kindly request you to grant us permission to collect the required image dataset under your supervision."
    )

    sig_tbl = doc.add_table(rows=1, cols=2)
    sig_tbl.autofit = False

    c_left = sig_tbl.cell(0, 0)
    c_left.width = Inches(3.4)
    p_s = c_left.paragraphs[0]
    p_s.paragraph_format.line_spacing = 1.15
    p_s.paragraph_format.space_after = Pt(0)
    r_s = p_s.add_run(
        "Submitted by (Project Team):\n\n\n"
        "_____________________________\n"
        "Student Lead / Team Representative\n"
        "Project: SmartSight\n"
        "Branch / Dept.: CSE / AI\n"
        "Contact: _____________________"
    )
    r_s.font.size = Pt(9.5)

    c_right = sig_tbl.cell(0, 1)
    c_right.width = Inches(3.4)
    p_m = c_right.paragraphs[0]
    p_m.paragraph_format.line_spacing = 1.15
    p_m.paragraph_format.space_after = Pt(0)
    r_m = p_m.add_run(
        "Permission Granted & Approved by:\n\n\n"
        "_____________________________\n"
        "Signature of Project Mentor / Guide\n"
        "Name: _______________________\n"
        "Designation: _________________\n"
        "Seal / Stamp: ________________"
    )
    r_m.font.size = Pt(9.5)

    doc.save(filename)
    print(f"Generated Mentor Permission Letter: {filename}")


if __name__ == "__main__":
    build_institutional_permission_doc()
    build_student_collection_sheet("s:/BKP/smartsight/Student_Dataset_Collection_Sheet.docx")
    try:
        build_student_collection_sheet("s:/BKP/smartsight/Student_Participation_and_Consent_Letter.docx")
    except Exception as e:
        pass
    build_mentor_permission_letter()


